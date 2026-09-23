"""One-off report generator: bonus category performance (last 30 days) and
a per-user bonus breakdown with each user's bonus total as a % of their
lifetime deposit (total_recharge). Produces an .xlsx with two sheets,
committed into the repo so it can be pulled and handed to the user. Always
writes debug/gen_bonus_performance_report.json (including a traceback on
failure) and commits both, since GitHub Actions logs for this repo aren't
readable without signing in. One-off; delete this script and its workflow
after use.
"""
import json
import os
import sqlite3
import subprocess
import traceback
from datetime import datetime, timedelta

import boto3
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
MASTER_DB = os.path.join(BASE, "master_userlist.db")
DAILY_DB = os.path.join(BASE, "daily_records.db")
OUT_XLSX = os.path.join(BASE, "debug", "bonus_performance_report.xlsx")

WINDOW_DAYS = 30


def r2_client():
    return boto3.client(
        "s3",
        endpoint_url=os.environ["R2_ENDPOINT_URL"],
        aws_access_key_id=os.environ["R2_ACCESS_KEY_ID"],
        aws_secret_access_key=os.environ["R2_SECRET_ACCESS_KEY"],
        region_name="auto",
    )


def run(result):
    bucket = os.environ["R2_BUCKET"]
    s3 = r2_client()
    s3.download_file(bucket, "master_userlist.db", MASTER_DB)
    s3.download_file(bucket, "daily_records.db", DAILY_DB)

    dconn = sqlite3.connect(DAILY_DB)
    dcur = dconn.cursor()

    window_start = (datetime.utcnow() + timedelta(hours=5, minutes=30) - timedelta(days=WINDOW_DAYS)).strftime("%Y-%m-%d")
    result["window_start"] = window_start

    bonus_rows = dcur.execute(
        "SELECT user_id, matched_category, change_value FROM bonuses "
        "WHERE matched_category IS NOT NULL AND create_time >= ?",
        (window_start,),
    ).fetchall()
    dconn.close()
    result["bonus_rows_scanned"] = len(bonus_rows)

    # Category performance
    from collections import defaultdict
    cat_stats = defaultdict(lambda: {"users": set(), "total_value": 0.0})
    user_stats = defaultdict(lambda: {"total_bonus": 0.0, "categories": defaultdict(float)})
    grand_total = 0.0
    for user_id, category, change_value in bonus_rows:
        value = change_value or 0.0
        grand_total += value
        c = cat_stats[category]
        c["users"].add(user_id)
        c["total_value"] += value
        if user_id is not None:
            u = user_stats[user_id]
            u["total_bonus"] += value
            u["categories"][category] += value

    category_rows = []
    for category, s in cat_stats.items():
        claimed_users = len(s["users"])
        category_rows.append({
            "category": category,
            "claimed_users": claimed_users,
            "total_value": round(s["total_value"], 2),
            "avg_per_user": round(s["total_value"] / claimed_users, 2) if claimed_users else 0.0,
            "pct_of_total_bonus": round(s["total_value"] / grand_total * 100, 2) if grand_total else 0.0,
        })
    category_rows.sort(key=lambda r: -r["total_value"])
    result["category_count"] = len(category_rows)
    result["grand_total_bonus"] = round(grand_total, 2)

    # Per-user report: total deposit (lifetime total_recharge) + agent
    mconn = sqlite3.connect(MASTER_DB)
    mcur = mconn.cursor()
    agent_by_user = {}
    try:
        agent_by_user = dict(mcur.execute("SELECT user_id, agent_name FROM agent_assignments").fetchall())
    except sqlite3.OperationalError:
        pass

    user_ids = list(user_stats.keys())
    total_recharge_by_user = {}
    CHUNK = 500
    for i in range(0, len(user_ids), CHUNK):
        chunk = user_ids[i:i + CHUNK]
        placeholders = ",".join("?" * len(chunk))
        for uid, total_recharge in mcur.execute(
            f"SELECT user_id, total_recharge FROM users WHERE user_id IN ({placeholders})", chunk
        ).fetchall():
            total_recharge_by_user[uid] = total_recharge or 0.0
    mconn.close()

    user_rows = []
    for user_id, s in user_stats.items():
        total_recharge = total_recharge_by_user.get(user_id, 0.0)
        top_category = max(s["categories"].items(), key=lambda kv: kv[1])[0] if s["categories"] else None
        user_rows.append({
            "user_id": user_id,
            "agent": agent_by_user.get(user_id) or "Un-Assigned",
            "total_bonus_claimed": round(s["total_bonus"], 2),
            "top_category": top_category,
            "category_count": len(s["categories"]),
            "total_deposit_lifetime": round(total_recharge, 2),
            "bonus_pct_of_deposit": round(s["total_bonus"] / total_recharge * 100, 2) if total_recharge else None,
        })
    user_rows.sort(key=lambda r: -r["total_bonus_claimed"])
    result["user_count"] = len(user_rows)

    # Write Excel
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Bonus Category Performance"
    ws1.append(["Bonus Category", "Claimed Users", "Total Value", "Avg Per User", "% of Total Bonus Paid"])
    for r in category_rows:
        ws1.append([r["category"], r["claimed_users"], r["total_value"], r["avg_per_user"], r["pct_of_total_bonus"]])

    ws2 = wb.create_sheet("Per-User Bonus Report")
    ws2.append([
        "User ID", "Agent", "Total Bonus Claimed", "Top Bonus Category", "Distinct Categories Claimed",
        "Total Deposit (Lifetime)", "Bonus % of Deposit",
    ])
    for r in user_rows:
        ws2.append([
            r["user_id"], r["agent"], r["total_bonus_claimed"], r["top_category"], r["category_count"],
            r["total_deposit_lifetime"], r["bonus_pct_of_deposit"],
        ])

    wb.save(OUT_XLSX)
    result["xlsx_path"] = OUT_XLSX
    result["status"] = "success"


def main():
    result = {}
    try:
        run(result)
    except Exception:
        result["status"] = "error"
        result["traceback"] = traceback.format_exc()
        print(result["traceback"])

    out_path = os.path.join(BASE, "debug")
    os.makedirs(out_path, exist_ok=True)
    with open(os.path.join(out_path, "gen_bonus_performance_report.json"), "w") as f:
        json.dump({k: v for k, v in result.items() if k != "xlsx_path"}, f, indent=2, default=str)

    subprocess.run(["git", "config", "user.email", "pipeline@bot.local"], check=True)
    subprocess.run(["git", "config", "user.name", "pipeline-bot"], check=True)
    subprocess.run(["git", "add", "-f", "debug/gen_bonus_performance_report.json"], check=True)
    if result.get("status") == "success":
        subprocess.run(["git", "add", "-f", OUT_XLSX], check=True)
    commit = subprocess.run(["git", "commit", "-m", "debug: bonus performance report"])
    if commit.returncode == 0:
        for attempt in range(5):
            push = subprocess.run(["git", "push"])
            if push.returncode == 0:
                break
            subprocess.run(["git", "pull", "--rebase", "origin", "main"], check=True)
        else:
            raise RuntimeError("git push failed after 5 rebase retries")
    print(json.dumps({k: v for k, v in result.items() if k != "xlsx_path"}, indent=2, default=str))

    if result.get("status") != "success":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
