import sqlite3
import openpyxl
import os

DB_PATH = "master_userlist.db"
if os.path.exists(DB_PATH):
    os.remove(DB_PATH)

conn = sqlite3.connect(DB_PATH)
cur = conn.cursor()

# ---- VIP level reference table (from attached image) ----
cur.execute("""
CREATE TABLE vip_levels (
    vip_level INTEGER PRIMARY KEY,
    member_name TEXT,
    upgrade_required_experience INTEGER,
    level_bonus_amount REAL,
    daily_withdrawal_times INTEGER,
    gift_scratch_card_name TEXT,
    gift_scratch_card_count INTEGER,
    spin_times INTEGER,
    withdrawal_fee_coefficient REAL,
    withdrawal_single_limit_amount REAL,
    status TEXT
)
""")

vip_rows = [
    (0, "V0", 0, 0, 0, "SAPPHIRE", 0, 0, 0, 0, "Active"),
    (1, "V1", 200, 0, 2, "Ice Cold", 0, 1, 0.035, 1000, "Active"),
    (2, "V2", 1500, 20, 2, "Ice Cold", 1, 1, 0.035, 2000, "Active"),
    (3, "V3", 9600, 50, 3, "Ice Cold", 2, 2, 0.035, 3000, "Active"),
    (4, "V4", 19600, 100, 3, "High Roller", 2, 3, 0.03, 5000, "Active"),
    (5, "V5", 95600, 200, 3, "High Roller", 2, 4, 0.03, 10000, "Active"),
    (6, "V6", 295600, 500, 3, "High Roller", 3, 5, 0.03, 20000, "Active"),
    (7, "V7", 795600, 800, 4, "POWER", 3, 7, 0.03, 20000, "Active"),
    (8, "V8", 1795600, 1200, 3, "POWER", 4, 10, 0.025, 50000, "Active"),
    (9, "V9", 3795600, 2000, 4, "POWER", 4, 15, 0.025, 50000, "Active"),
    (10, "V10", 8795600, 3500, 4, "POWER", 12, 3, 0.02, 50000, "Active"),
    (11, "V11", 16795600, 5000, 4, "POWER", 10, 12, 0.02, 50000, "Active"),
    (12, "V12", 28795600, 7000, 4, "POWER", 14, 13, 0.015, 50000, "Active"),
    (13, "V13", 44795600, 9000, 5, "POWER", 4, 4, 0.015, 50000, "Active"),
    (14, "V14", 69795600, 12000, 5, "POWER", 5, 4, 0.015, 50000, "Active"),
    (15, "V15", 119795600, 15000, 5, "POWER", 100, 100, 0.01, 50000, "Active"),
]
cur.executemany("INSERT INTO vip_levels VALUES (?,?,?,?,?,?,?,?,?,?,?)", vip_rows)

# ---- Users table ----
# Chinese header -> English column mapping (lotteryUserInfo files)
col_map = [
    ("user_id", "INTEGER"),
    ("agent_status", "INTEGER"),
    ("agent_user_id", "TEXT"),
    ("parent_user_id", "TEXT"),
    ("direct_parent", "TEXT"),
    ("agent_level1", "TEXT"),
    ("agent_level2", "TEXT"),
    ("agent_level3", "TEXT"),
    ("agent_level4", "TEXT"),
    ("agent_level", "TEXT"),
    ("username", "TEXT"),
    ("gender", "INTEGER"),
    ("phone", "TEXT"),
    ("email", "TEXT"),
    ("register_ip", "TEXT"),
    ("birth_date", "TEXT"),
    ("app_version", "TEXT"),
    ("register_device", "TEXT"),
    ("login_device", "TEXT"),
    ("register_channel", "TEXT"),
    ("is_test_account", "INTEGER"),
    ("invited_by_user_id", "TEXT"),
    ("register_source", "TEXT"),
    ("last_active_time", "TEXT"),
    ("last_login_device", "TEXT"),
    ("device_id", "TEXT"),
    ("user_status", "INTEGER"),
    ("push_token", "TEXT"),
    ("vip_level", "INTEGER"),
    ("register_app_version", "TEXT"),
    ("channel", "TEXT"),
    ("balance", "REAL"),
    ("recharge_count", "INTEGER"),
    ("query_time", "TEXT"),
    ("start_time", "TEXT"),
    ("end_time", "TEXT"),
    ("recharge_count_start", "TEXT"),
    ("recharge_count_end", "TEXT"),
    ("user_balance", "REAL"),
    ("total_recharge", "REAL"),
    ("frozen_amount", "REAL"),
    ("total_withdrawal", "REAL"),
    ("withdrawal_quota", "REAL"),
    ("city", "TEXT"),
    ("mark", "TEXT"),
    ("followup_time", "TEXT"),
    ("next_followup_time", "TEXT"),
    ("tag", "TEXT"),
    ("im_user_id", "TEXT"),
    ("im_user_status", "TEXT"),
    ("group_name", "TEXT"),
    ("adjust_adid", "TEXT"),
    ("im_customer", "TEXT"),
    ("create_time", "TEXT"),
    ("update_time", "TEXT"),
    ("package_id", "REAL"),
]

cols_sql = ",\n    ".join(f"{name} {typ}" for name, typ in col_map[1:])
cur.execute(f"""
CREATE TABLE users (
    user_id INTEGER PRIMARY KEY,
    {cols_sql}
)
""")
cur.execute("CREATE INDEX idx_users_phone ON users(phone)")
cur.execute("CREATE INDEX idx_users_vip ON users(vip_level)")
cur.execute("CREATE INDEX idx_users_update_time ON users(update_time)")

conn.commit()

# ---- Load and merge both lotteryUserInfo files, dedupe by user_id keeping latest update_time ----
files = [
    "/Users/devtr/Downloads/lotteryUserInfo_1782975504094.xlsx",
    "/Users/devtr/Downloads/lotteryUserInfo_1782975702238.xlsx",
]

placeholders = ",".join(["?"] * len(col_map))
insert_sql = f"INSERT INTO users VALUES ({placeholders})"
update_time_idx = len(col_map) - 2  # update_time is second to last

best_rows = {}  # user_id -> row tuple
total_seen = 0

for f in files:
    wb = openpyxl.load_workbook(f, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(min_row=2, values_only=True)
    for row in rows_iter:
        if row[0] is None:
            continue
        total_seen += 1
        uid = int(row[0])
        row = list(row)
        row[0] = uid
        row = tuple(row)
        ut = row[update_time_idx]
        existing = best_rows.get(uid)
        if existing is None:
            best_rows[uid] = row
        else:
            existing_ut = existing[update_time_idx]
            try:
                if ut is not None and (existing_ut is None or str(ut) > str(existing_ut)):
                    best_rows[uid] = row
            except Exception:
                best_rows[uid] = row
    wb.close()

print(f"Total rows scanned: {total_seen}, unique users: {len(best_rows)}")

def clean(row):
    return tuple(str(v) if hasattr(v, "isoformat") else v for v in row)

batch = [clean(r) for r in best_rows.values()]
cur.executemany(insert_sql, batch)
conn.commit()

cur.execute("SELECT COUNT(*) FROM users")
print("users table rows:", cur.fetchone()[0])

conn.close()
print("Done ->", DB_PATH)
